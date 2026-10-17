"""

DataLens forensic acquisition and AI triage engine.

Supports:

- Read-only discovery of mounted removable storage (Windows/Linux)

- Directory and uploaded-file acquisition into the evidence vault

- SHA-256 verification through evidence_manager

- Transparent rule-based triage

- Random Forest classification and Isolation Forest anomaly detection

- Case-level analysis saved against individual evidence records

Important:

AI outputs are triage indicators, not proof that a file is malicious.

A mounted USB device is not automatically write-blocked or bit-for-bit imaged.

"""

import json

import math

import os

import platform

import re

import subprocess

import tempfile

from pathlib import Path

import psutil

# Machine-learning libraries are optional at application startup. On some Windows

# systems, Application Control may block NumPy's compiled extension. In that case,

# DataLens should still start and offer rule-based triage instead of crashing.

SKLEARN_AVAILABLE = False

SKLEARN_IMPORT_ERROR = None

try:

    from sklearn.ensemble import IsolationForest, RandomForestClassifier

    from sklearn.impute import SimpleImputer

    from sklearn.pipeline import Pipeline

    from sklearn.preprocessing import StandardScaler

    SKLEARN_AVAILABLE = True

except Exception as exc:  # Includes blocked/missing NumPy or scikit-learn binaries.

    IsolationForest = None

    RandomForestClassifier = None

    SimpleImputer = None

    Pipeline = None

    StandardScaler = None

    SKLEARN_IMPORT_ERROR = f"{type(exc).__name__}: {exc}"

from evidence_manager import (

    DATA_DIR,

    VAULT_DIR,

    get_evidence,

    register_evidence,

    save_analysis_result,

    verify_evidence,

)

RISK_EXTENSIONS = {

    ".exe", ".dll", ".scr", ".bat", ".cmd", ".ps1", ".vbs",

    ".js", ".jar", ".msi", ".hta", ".com", ".sh", ".pif",

}

SCRIPT_EXTENSIONS = {

    ".ps1", ".bat", ".cmd", ".vbs", ".js", ".sh", ".hta", ".py",

}

TEXT_EXTENSIONS = {

    ".txt", ".csv", ".json", ".xml", ".log", ".py", ".ps1",

    ".sh", ".html", ".htm", ".yaml", ".yml", ".ini", ".conf",

}

SENSITIVE_PATTERNS = {

    "Possible email address": r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",

    "Possible private key marker": r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----",

    "Possible credential assignment": r"(?i)\b(?:password|passwd|secret|api_key)\s*[:=]\s*\S+",

}

MAX_TEXT_SCAN_BYTES = 2 * 1024 * 1024
MAX_EXTRACTED_TEXT_CHARS = 2 * 1024 * 1024
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".doc"}

MAX_ENTROPY_BYTES = 32 * 1024 * 1024

MAX_FILES_PER_ACQUISITION = 10000

def is_inside(path, directory):

    """Return True if path resolves inside directory."""

    try:

        Path(path).resolve().relative_to(Path(directory).resolve())

        return True

    except (ValueError, OSError, RuntimeError):

        return False

def format_device_size(size_bytes):

    """Convert bytes to a readable size."""

    size = float(size_bytes or 0)

    for unit in ("B", "KB", "MB", "GB", "TB"):

        if size < 1024 or unit == "TB":

            return f"{size:.2f} {unit}"

        size /= 1024

    return f"{size:.2f} TB"

def file_entropy(file_path, max_bytes=MAX_ENTROPY_BYTES):

    """

    Calculate Shannon entropy from up to max_bytes.

    A high value may indicate compression/encryption; it is not proof of malware.

    """

    counts = [0] * 256

    total = 0

    with open(file_path, "rb") as file:

        remaining = max_bytes

        while remaining > 0:

            chunk = file.read(min(1024 * 1024, remaining))

            if not chunk:

                break

            remaining -= len(chunk)

            total += len(chunk)

            for value in chunk:

                counts[value] += 1

    if total == 0:

        return 0.0

    return -sum(

        (count / total) * math.log2(count / total)

        for count in counts if count

    )

def _get_windows_usb_devices():

    """List Windows logical drives reported by Windows as removable."""

    script = r"""

    $items = @(

      Get-CimInstance Win32_LogicalDisk |

      Where-Object { $_.DriveType -eq 2 } |

      ForEach-Object {

        [PSCustomObject]@{

          name = $_.VolumeName

          mountpoint = $_.DeviceID

          filesystem = $_.FileSystem

          size_bytes = $_.Size

          free_bytes = $_.FreeSpace

          drive_type = "Removable"

        }

      }

    )

    ConvertTo-Json -InputObject $items -Compress

    """

    try:

        result = subprocess.run(

            ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],

            capture_output=True, text=True, timeout=15, check=True,

        )

        raw = result.stdout.strip()

        if not raw:

            return []

        data = json.loads(raw)

        if isinstance(data, dict):

            data = [data]

        devices = []

        for item in data:

            mountpoint = item.get("mountpoint")

            if not mountpoint:

                continue

            devices.append({

                "name": item.get("name") or "Removable drive",

                "mountpoint": mountpoint,

                "filesystem": item.get("filesystem") or "Unknown",

                "size_bytes": int(item.get("size_bytes") or 0),

                "free_bytes": int(item.get("free_bytes") or 0),

                "drive_type": "Removable",

                "platform": "Windows",

            })

        return devices

    except (subprocess.SubprocessError, OSError, ValueError, json.JSONDecodeError):

        return []

def _get_linux_removable_devices():

    """List mounted USB/removable partitions on Linux using lsblk."""

    try:

        result = subprocess.run(

            ["lsblk", "-J", "-o", "NAME,PATH,TYPE,RM,TRAN,LABEL,FSTYPE,SIZE,MOUNTPOINTS"],

            capture_output=True, text=True, timeout=15, check=True,

        )

        data = json.loads(result.stdout)

        devices = []

        def visit(node, parent_removable=False):

            removable = (

                parent_removable

                or str(node.get("rm", "0")) == "1"

                or str(node.get("tran", "")).lower() == "usb"

            )

            mountpoints = node.get("mountpoints") or []

            if isinstance(mountpoints, str):

                mountpoints = [mountpoints]

            if removable:

                for mountpoint in mountpoints:

                    if not mountpoint:

                        continue

                    try:

                        usage = psutil.disk_usage(mountpoint)

                    except (OSError, PermissionError):

                        continue

                    devices.append({

                        "name": node.get("label") or node.get("name") or "USB device",

                        "mountpoint": mountpoint,

                        "filesystem": node.get("fstype") or "Unknown",

                        "size_bytes": int(usage.total),

                        "free_bytes": int(usage.free),

                        "drive_type": "USB/removable",

                        "platform": "Linux",

                    })

            for child in node.get("children") or []:

                visit(child, removable)

        for device in data.get("blockdevices", []):

            visit(device)

        return devices

    except (subprocess.SubprocessError, OSError, ValueError, json.JSONDecodeError):

        return []

def detect_attached_devices():

    """

    Discover mounted removable devices on Windows or Linux.

    This is discovery only. It does not provide write blocking or

    guarantee forensic soundness of a mounted-file acquisition.

    """

    system = platform.system().lower()

    if system == "windows":

        return _get_windows_usb_devices()

    if system == "linux":

        return _get_linux_removable_devices()

    return []

def enumerate_files(source_directory, max_files=MAX_FILES_PER_ACQUISITION):

    """

    Enumerate regular files under an authorized source directory.

    Skips symbolic links and excludes the DataLens vault itself.

    """

    source = Path(source_directory).expanduser().resolve(strict=True)

    if not source.is_dir():

        raise ValueError("The selected source must be a directory.")

    if is_inside(source, VAULT_DIR):

        raise ValueError("The evidence vault cannot be used as its own source.")

    if is_inside(VAULT_DIR, source):

        raise ValueError(

            "The selected source contains the DataLens vault. Select a narrower directory."

        )

    discovered = []

    for root, dirs, files in os.walk(source, followlinks=False):

        root_path = Path(root)

        dirs[:] = [

            name for name in dirs

            if not (root_path / name).is_symlink()

        ]

        for name in files:

            candidate = root_path / name

            try:

                if candidate.is_symlink() or not candidate.is_file():

                    continue

                discovered.append(candidate)

                if len(discovered) >= max_files:

                    return source, discovered

            except OSError:

                continue

    return source, discovered

def acquire_directory(source_directory, case_id, actor, max_files=MAX_FILES_PER_ACQUISITION):

    """Acquire each file from a directory and register it as case evidence."""

    source, files = enumerate_files(source_directory, max_files=max_files)

    results = []

    for source_file in files:

        try:

            relative_name = source_file.relative_to(source).as_posix()

            record = register_evidence(

                case_id=case_id,

                source_file=source_file,

                actor=actor,

                relative_name=relative_name,

            )

            results.append({

                "Evidence ID": record["evidence_id"],

                "File": relative_name,

                "Size (bytes)": record.get("size_bytes", source_file.stat().st_size),

                "SHA-256": record.get("sha256", ""),

                "Status": record.get("status", "Registered"),

            })

        except Exception as exc:

            try:

                display_name = source_file.relative_to(source).as_posix()

            except ValueError:

                display_name = source_file.name

            results.append({

                "Evidence ID": "-",

                "File": display_name,

                "Size (bytes)": "-",

                "SHA-256": "-",

                "Status": f"FAILED: {exc}",

            })

    successful = sum(

        1 for row in results

        if str(row.get("Status", "")).lower() in {

            "acquired and hash-verified", "registered", "acquired"

        }

    )

    return {

        "total": len(results),

        "successful": successful,

        "failed": len(results) - successful,

        "limit_reached": len(files) >= max_files,

        "results": results,

    }

def acquire_uploaded_file(uploaded_file, case_id, actor):

    """Acquire a Streamlit uploaded file into the evidence vault."""

    original_name = Path(uploaded_file.name).name

    if not original_name:

        raise ValueError("The uploaded file has no valid filename.")

    temporary_path = None

    try:

        DATA_DIR.mkdir(parents=True, exist_ok=True)

        with tempfile.NamedTemporaryFile(

            prefix="datalens_upload_", dir=DATA_DIR, delete=False

        ) as temporary:

            temporary_path = Path(temporary.name)

            uploaded_file.seek(0)

            while True:

                chunk = uploaded_file.read(1024 * 1024)

                if not chunk:

                    break

                temporary.write(chunk)

        return register_evidence(

            case_id=case_id,

            source_file=temporary_path,

            actor=actor,

            relative_name=original_name,

        )

    finally:

        if temporary_path:

            temporary_path.unlink(missing_ok=True)

def acquire_selected_files(file_paths, case_id, actor):

    """Acquire a sequence of file paths into a case."""

    results = []

    for file_path in file_paths:

        try:

            record = register_evidence(

                case_id=case_id,

                source_file=file_path,

                actor=actor,

            )

            results.append({

                "Evidence ID": record["evidence_id"],

                "File": record.get("file_name", Path(file_path).name),

                "Status": record.get("status", "Registered"),

                "SHA-256": record.get("sha256", ""),

            })

        except Exception as exc:

            results.append({

                "Evidence ID": "-",

                "File": str(file_path),

                "Status": f"FAILED: {exc}",

                "SHA-256": "-",

            })

    return results

def _extract_features(evidence):

    """

    Extract numeric features for ML triage.

    No evidence file is executed. Content sampling is bounded.

    """

    path = Path(evidence["stored_path"])

    stat = path.stat()

    suffix = path.suffix.lower()

    filename = path.name.lower()

    size = stat.st_size

    entropy = 0.0

    if path.is_file() and size:

        try:

            entropy = file_entropy(path)

        except OSError:

            entropy = 0.0

    return {

        "size_bytes": float(size),

        "log_size": math.log1p(size),

        "entropy": float(entropy),

        "extension_is_executable": float(suffix in RISK_EXTENSIONS),

        "extension_is_script": float(suffix in SCRIPT_EXTENSIONS),

        "has_double_extension": float(bool(re.search(

            r"\.(?:pdf|docx?|xlsx?|jpg|jpeg|png|txt)\.(?:exe|scr|bat|cmd|js|vbs|ps1)$",

            filename,

        ))),

        "extension_length": float(len(suffix)),

        "filename_length": float(len(filename)),

        "is_hidden_name": float(filename.startswith(".")),

        "has_space_in_name": float(" " in filename),

    }

def extract_document_text(file_path):
    """Extract bounded text from PDF/DOCX, or legacy DOC when a CLI parser exists.

    Optional dependencies are used only for extraction. Evidence files are opened
    read-only and are never converted in place or executed.
    """
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF extraction requires pypdf. Install it with: python -m pip install pypdf") from exc

        reader = PdfReader(str(path), strict=False)
        chunks = []
        total = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            remaining = MAX_EXTRACTED_TEXT_CHARS - total
            if remaining <= 0:
                break
            chunks.append(page_text[:remaining])
            total += min(len(page_text), remaining)
        return "\n".join(chunks)

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX extraction requires python-docx. Install it with: python -m pip install python-docx") from exc

        document = Document(str(path))
        chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(chunks)[:MAX_EXTRACTED_TEXT_CHARS]

    if suffix == ".doc":
        # Legacy .doc is binary. Use a dedicated local parser if installed.
        for command in ("antiword", "catdoc"):
            try:
                result = subprocess.run(
                    [command, str(path)], capture_output=True, timeout=20, check=False
                )
            except (OSError, subprocess.TimeoutExpired):
                continue
            if result.returncode == 0 and result.stdout:
                return result.stdout[:MAX_EXTRACTED_TEXT_CHARS].decode("utf-8", errors="replace")
        raise RuntimeError(
            "Legacy .doc extraction needs antiword or catdoc installed and available on PATH. "
            "The file was not decoded as plain text."
        )

    raise ValueError(f"Unsupported document format: {suffix or '(no extension)'}")


def inspect_evidence(evidence):

    """

    Rule-based triage of one acquired file.

    This is a heuristic review aid, not a definitive malware verdict.

    """

    path = Path(evidence["stored_path"])

    if not path.is_file():

        return {

            "risk_score": 100,

            "severity": "CRITICAL",

            "integrity": "MISSING",

            "findings": ["Evidence file is missing from the vault."],

            "entropy": None,

            "analysis_method": "Rule-based forensic triage",

        }

    integrity_result = verify_evidence(evidence["evidence_id"])

    integrity = integrity_result.get("status", "UNKNOWN")

    findings = []

    score = 0

    suffix = path.suffix.lower()

    filename = path.name.lower()

    if integrity == "MISMATCH":

        findings.append("SHA-256 mismatch: stored evidence differs from its acquisition hash.")

        score += 100

    elif integrity == "MISSING":

        findings.append("Evidence is missing from the vault.")

        score += 100

    if suffix in RISK_EXTENSIONS:

        findings.append("Executable or potentially risky extension; review provenance and context.")

        score += 25

    if suffix in SCRIPT_EXTENSIONS:

        findings.append("Script file detected; inspect its contents and origin.")

        score += 10

    if re.search(

        r"\.(?:pdf|docx?|xlsx?|jpg|jpeg|png|txt)\.(?:exe|scr|bat|cmd|js|vbs|ps1)$",

        filename,

    ):

        findings.append("Potentially misleading double extension detected.")

        score += 35

    size = path.stat().st_size

    entropy = None

    # Scan plain-text files as UTF-8 and document formats through dedicated
    # extractors. Never decode arbitrary binary documents as plain text.
    if size <= MAX_TEXT_SCAN_BYTES and (suffix in TEXT_EXTENSIONS or suffix in DOCUMENT_EXTENSIONS):
        try:
            if suffix in TEXT_EXTENSIONS:
                raw = path.read_bytes()
                text_content = "" if b"\x00" in raw else raw.decode("utf-8", errors="replace")
            else:
                text_content = extract_document_text(path)

            for label, pattern in SENSITIVE_PATTERNS.items():
                if re.search(pattern, text_content):
                    findings.append(label + " detected in extracted content; inspect securely.")
                    score += 10
            entropy = file_entropy(path)
        except Exception as exc:
            # Parser errors vary by optional library; record them without stopping
            # analysis of the rest of the case. No evidence is changed.
            findings.append(f"Content inspection unavailable: {exc}")
    elif suffix in DOCUMENT_EXTENSIONS:
        findings.append(
            f"Document content was not scanned because the file exceeds the "
            f"{MAX_TEXT_SCAN_BYTES // (1024 * 1024)} MiB inspection limit."
        )

    if entropy is not None and entropy > 7.5:

        findings.append(

            "High entropy observed; compression, encryption, or encoding may explain it."

        )

        score += 10

    if not findings:

        findings.append("No indicators from the configured checks were detected.")

    score = min(score, 100)

    if integrity != "VERIFIED":

        severity = "CRITICAL"

    elif score >= 60:

        severity = "HIGH"

    elif score >= 30:

        severity = "MEDIUM"

    elif score > 0:

        severity = "LOW"

    else:

        severity = "NO INDICATOR"

    return {

        "risk_score": score,

        "severity": severity,

        "integrity": integrity,

        "findings": findings,

        "entropy": round(entropy, 3) if entropy is not None else None,

        "analysis_method": "Rule-based forensic triage",

        "interpretation": (

            "Indicators support prioritization and require analyst review; "

            "they do not prove malicious activity."

        ),

    }

def _demo_labels_from_features(feature_rows):

    """Return demonstration labels derived from transparent heuristics."""

    labels = []

    for row in feature_rows:

        heuristic_score = (

            25 * row["extension_is_executable"]

            + 10 * row["extension_is_script"]

            + 35 * row["has_double_extension"]

            + 8 * float(row["entropy"] >= 7.5)

        )

        labels.append("review" if heuristic_score >= 25 else "baseline")

    # Do not fabricate a second class to make the demo appear more capable.

    # These are heuristic-derived labels, not independently verified ground truth.

    return labels

def _train_random_forest_for_demo(feature_rows):

    """Fit the sklearn demonstration forest when its native dependencies load."""

    feature_names = list(feature_rows[0].keys())

    X = [[row[name] for name in feature_names] for row in feature_rows]

    labels = _demo_labels_from_features(feature_rows)

    model = Pipeline([

        ("imputer", SimpleImputer(strategy="constant", fill_value=0)),

        ("scaler", StandardScaler()),

        ("classifier", RandomForestClassifier(

            n_estimators=100, random_state=42, class_weight="balanced"

        )),

    ])

    model.fit(X, labels)

    return model, feature_names, X

def _gini_from_counts(counts, total):

    if total <= 0:

        return 0.0

    return 1.0 - sum((count / total) ** 2 for count in counts.values())

def _pure_python_random_forest(feature_rows, seed=42, n_estimators=25):

    """Small, dependency-free Random Forest demo used if sklearn is blocked.

    It uses bootstrap samples, random feature subsets, and Gini-based splits.

    Training labels are still generated by heuristics, not independently

    labeled evidence. Its vote fraction is an uncalibrated agreement measure.

    Training is capped to 500 records to keep the fallback responsive.

    """

    import random

    from collections import Counter

    if not feature_rows:

        return [], []

    feature_names = list(feature_rows[0].keys())

    X = [[float(row.get(name, 0) or 0) for name in feature_names] for row in feature_rows]

    all_labels = _demo_labels_from_features(feature_rows)

    n_samples = len(X)

    rng = random.Random(seed)

    # Bound the pure-Python training cost for larger cases. Keep one sample of

    # each generated class in the pool when both classes exist.

    if n_samples > 500:

        pool = rng.sample(range(n_samples), 500)

        for label in set(all_labels):

            if not any(all_labels[i] == label for i in pool):

                candidate = next(i for i, value in enumerate(all_labels) if value == label)

                pool[-1] = candidate

        pool = sorted(set(pool))

    else:

        pool = list(range(n_samples))

    train_X = [X[i] for i in pool]

    train_y = [all_labels[i] for i in pool]

    n_features = len(feature_names)

    feature_subset_size = max(1, int(math.sqrt(max(1, n_features))))

    max_depth = max(1, min(8, math.ceil(math.log2(max(2, len(pool)))) + 1))

    def majority_label(labels):

        counts = Counter(labels)

        return sorted(counts, key=lambda label: (-counts[label], label))[0]

    def find_best_split(indices, candidate_features):

        labels_here = [train_y[index] for index in indices]

        parent_counts = Counter(labels_here)

        parent_impurity = _gini_from_counts(parent_counts, len(indices))

        best_gain = 0.0

        best_split = None

        for feature_index in candidate_features:

            ordered = sorted((train_X[index][feature_index], train_y[index]) for index in indices)

            left_counts = Counter()

            right_counts = Counter(label for _, label in ordered)

            for position in range(len(ordered) - 1):

                value, label = ordered[position]

                left_counts[label] += 1

                right_counts[label] -= 1

                if right_counts[label] == 0:

                    del right_counts[label]

                next_value = ordered[position + 1][0]

                if value == next_value:

                    continue

                left_n = position + 1

                right_n = len(ordered) - left_n

                weighted = (

                    (left_n / len(ordered)) * _gini_from_counts(left_counts, left_n)

                    + (right_n / len(ordered)) * _gini_from_counts(right_counts, right_n)

                )

                gain = parent_impurity - weighted

                if gain > best_gain:

                    best_gain = gain

                    best_split = (feature_index, (value + next_value) / 2.0)

        return best_split

    def build_tree(indices, depth, tree_rng):

        local_labels = [train_y[index] for index in indices]

        majority = majority_label(local_labels)

        if depth >= max_depth or len(indices) < 2 or len(set(local_labels)) < 2:

            return {"leaf": True, "prediction": majority}

        candidate_features = tree_rng.sample(

            range(n_features), min(feature_subset_size, n_features)

        )

        split = find_best_split(indices, candidate_features)

        if split is None:

            return {"leaf": True, "prediction": majority}

        feature_index, threshold = split

        left_indices = [i for i in indices if train_X[i][feature_index] <= threshold]

        right_indices = [i for i in indices if train_X[i][feature_index] > threshold]

        if not left_indices or not right_indices:

            return {"leaf": True, "prediction": majority}

        return {

            "leaf": False,

            "feature": feature_index,

            "threshold": threshold,

            "left": build_tree(left_indices, depth + 1, tree_rng),

            "right": build_tree(right_indices, depth + 1, tree_rng),

        }

    def predict_tree(tree, row):

        node = tree

        while not node["leaf"]:

            node = node["left"] if row[node["feature"]] <= node["threshold"] else node["right"]

        return node["prediction"]

    trees = []

    pool_size = len(pool)

    for tree_index in range(n_estimators):

        tree_rng = random.Random(seed + tree_index + 1)

        bootstrap = [pool_index for pool_index in tree_rng.choices(range(pool_size), k=pool_size)]

        trees.append(build_tree(bootstrap, 0, tree_rng))

    predictions = []

    vote_fractions = []

    for row in X:

        votes = Counter(predict_tree(tree, row) for tree in trees)

        chosen = sorted(votes, key=lambda label: (-votes[label], label))[0]

        predictions.append(chosen)

        vote_fractions.append(votes[chosen] / max(1, len(trees)))

    return predictions, vote_fractions

def _expected_isolation_path_length(sample_count):

    """Expected path length correction used by the Isolation Forest score."""

    if sample_count <= 1:

        return 0.0

    if sample_count == 2:

        return 1.0

    harmonic = sum(1.0 / value for value in range(1, sample_count))

    return 2.0 * harmonic - (2.0 * (sample_count - 1) / sample_count)

def _pure_python_isolation_forest(feature_rows, seed=42, n_estimators=50):

    """Dependency-free Isolation Forest demo for small/medium case collections.

    This implements randomized isolation trees without NumPy/scikit-learn.

    Scores are exploratory and should not be treated as malware verdicts.

    """

    import random

    n_samples = len(feature_rows)

    if n_samples < 3:

        return ["Not enough case files"] * n_samples, [None] * n_samples

    feature_names = list(feature_rows[0].keys())

    X = [[float(row.get(name, 0) or 0) for name in feature_names] for row in feature_rows]

    rng = random.Random(seed)

    sample_size = min(256, n_samples)

    max_depth = max(1, math.ceil(math.log2(sample_size)))

    normalizer = _expected_isolation_path_length(sample_size)

    def build_tree(indices, depth, tree_rng):

        if len(indices) <= 1 or depth >= max_depth:

            return {"leaf": True, "size": len(indices)}

        splittable = []

        for feature_index in range(len(feature_names)):

            values = [X[index][feature_index] for index in indices]

            low, high = min(values), max(values)

            if low < high:

                splittable.append((feature_index, low, high))

        if not splittable:

            return {"leaf": True, "size": len(indices)}

        feature_index, low, high = tree_rng.choice(splittable)

        threshold = tree_rng.uniform(low, high)

        left_indices = [index for index in indices if X[index][feature_index] <= threshold]

        right_indices = [index for index in indices if X[index][feature_index] > threshold]

        if not left_indices or not right_indices:

            return {"leaf": True, "size": len(indices)}

        return {

            "leaf": False,

            "size": len(indices),

            "feature": feature_index,

            "threshold": threshold,

            "left": build_tree(left_indices, depth + 1, tree_rng),

            "right": build_tree(right_indices, depth + 1, tree_rng),

        }

    def path_length(tree, row, depth=0):

        if tree["leaf"]:

            return depth + _expected_isolation_path_length(tree["size"])

        branch = tree["left"] if row[tree["feature"]] <= tree["threshold"] else tree["right"]

        return path_length(branch, row, depth + 1)

    trees = []

    for tree_index in range(n_estimators):

        tree_rng = random.Random(seed + tree_index + 1001)

        indices = tree_rng.sample(range(n_samples), sample_size)

        trees.append(build_tree(indices, 0, tree_rng))

    labels = []

    decision_scores = []

    for row in X:

        average_path = sum(path_length(tree, row) for tree in trees) / len(trees)

        anomaly_score = 2.0 ** (-average_path / normalizer) if normalizer else 0.5

        # The classic isolation score treats values above 0.5 as more anomalous.

        labels.append("Anomaly" if anomaly_score > 0.5 else "Not flagged")

        decision_scores.append(round(0.5 - anomaly_score, 6))

    return labels, decision_scores

def analyze_case_ai(case_id, actor="Investigator"):

    """

    Run rule-based triage, Random Forest and Isolation Forest over all

    evidence in a case. Save a combined finding for every evidence record.

    Random Forest uses a heuristic-labeled demonstration fallback when no

    independently labeled training set has been configured. Its results must

    be treated as experimental and not as validated malware classification.

    """

    evidence_items = get_evidence(case_id)

    if not evidence_items:

        raise ValueError("No evidence has been acquired for this case.")

    usable_items = []

    feature_rows = []

    rule_findings = []

    skipped = []

    for evidence in evidence_items:

        try:

            path = Path(evidence["stored_path"])

            if not path.is_file():

                skipped.append({

                    "Evidence ID": evidence.get("evidence_id", ""),

                    "File": evidence.get("file_name", ""),

                    "Status": "Skipped: evidence file missing",

                })

                continue

            feature_rows.append(_extract_features(evidence))

            usable_items.append(evidence)

            rule_findings.append(inspect_evidence(evidence))

        except Exception as exc:

            skipped.append({

                "Evidence ID": evidence.get("evidence_id", ""),

                "File": evidence.get("file_name", ""),

                "Status": f"Skipped: {exc}",

            })

    if not usable_items:

        raise ValueError("No available evidence files could be analyzed.")

    if SKLEARN_AVAILABLE:

        rf_model, rf_feature_names, rf_X = _train_random_forest_for_demo(feature_rows)

        rf_predictions = [str(value) for value in rf_model.predict(rf_X)]

        # The model is trained on heuristic-derived labels from this same case

        # and then predicts those same records. In-sample probabilities would be

        # misleading, so do not report them as confidence values.

        rf_confidences = [None] * len(usable_items)

        isolation_labels = ["Not enough case files"] * len(usable_items)

        isolation_scores = [None] * len(usable_items)

        if len(usable_items) >= 3:

            feature_names = list(feature_rows[0].keys())

            X = [[row[name] for name in feature_names] for row in feature_rows]

            iso = Pipeline([

                ("imputer", SimpleImputer(strategy="constant", fill_value=0)),

                ("scaler", StandardScaler()),

                ("model", IsolationForest(

                    n_estimators=100,

                    contamination="auto",

                    random_state=42,

                )),

            ])

            iso.fit(X)

            predictions = iso.predict(X)

            scores = iso.decision_function(X)

            isolation_labels = [

                "Anomaly" if value == -1 else "Not flagged"

                for value in predictions

            ]

            isolation_scores = [round(float(value), 6) for value in scores]

        rf_note = (

            "Scikit-learn Random Forest trained on heuristic-derived demonstration labels; "

            "not validated against independent ground truth. Confidence is not a calibrated risk probability."

        )

        isolation_note = (

            "Scikit-learn Isolation Forest identifies files unusual relative to this case's feature distribution; "

            "an anomaly is not proof of maliciousness."

        )

        model_notes = [

            "Random Forest uses heuristic-derived demonstration labels; no independent ground-truth dataset is configured.",

            "RF confidence is N/A because in-sample probabilities or tree-vote fractions are not validated confidence estimates.",

            "Use an independently labeled dataset before interpreting Random Forest output as validated classification.",

            "Isolation Forest is a case-relative anomaly detector and may flag harmless unusual files.",

        ]

    else:

        # A small pure-Python implementation keeps demo triage available when

        # Windows Application Control blocks NumPy/scikit-learn binary modules.

        # It does not bypass or modify the operating-system security policy.

        rf_predictions, _vote_fractions = _pure_python_random_forest(feature_rows)

        # Tree vote fractions are not calibrated confidence and the labels are

        # generated from heuristics, so show N/A in the shared output schema.

        rf_confidences = [None] * len(usable_items)

        isolation_labels, isolation_scores = _pure_python_isolation_forest(feature_rows)

        rf_note = (

            "Dependency-free pure-Python demonstration Random Forest. Training labels are generated from "

            "heuristics, and the vote fraction is not calibrated. This is not equivalent to scikit-learn "

            "and is not validated against independent ground truth."

        )

        isolation_note = (

            "Dependency-free pure-Python Isolation Forest. Scores are exploratory and case-relative; "

            "an anomaly is not proof of maliciousness."

        )

        model_notes = [

            "NumPy/scikit-learn could not load, so DataLens used dependency-free pure-Python demo models.",

            "Random Forest labels remain heuristic-derived and are not validated classification.",

            "RF confidence is N/A because tree-vote fractions are not calibrated confidence estimates.",

            "Isolation Forest is case-relative and may flag harmless unusual files.",

            f"Original native dependency error: {SKLEARN_IMPORT_ERROR or 'unknown error'}",

        ]

    results = []

    for index, evidence in enumerate(usable_items):

        finding = dict(rule_findings[index])

        rf_confidence = rf_confidences[index]

        rf_prediction = str(rf_predictions[index])

        combined = {

            **finding,

            "ai_models": {

                "random_forest": {

                    "prediction": rf_prediction,

                    "confidence": round(rf_confidence, 4) if rf_confidence is not None else None,

                    "model_note": rf_note,

                },

                "isolation_forest": {

                    "result": isolation_labels[index],

                    "decision_score": isolation_scores[index],

                    "model_note": isolation_note,

                },

            },

            "features": feature_rows[index],

        }

        save_analysis_result(

            evidence_id=evidence["evidence_id"],

            findings=combined,

            actor=actor,

        )

        results.append({

            "Evidence ID": evidence["evidence_id"],

            "File": evidence.get("file_name", ""),

            "Rule severity": finding["severity"],

            "Rule risk score": finding["risk_score"],

            "Random Forest": rf_prediction,

            "RF confidence": (

                round(rf_confidence, 3) if rf_confidence is not None else "N/A"

            ),

            "Isolation Forest": isolation_labels[index],

            "Isolation score": isolation_scores[index],

            "Integrity": finding["integrity"],

        })

    return {

        "case_id": case_id,

        "total_evidence": len(evidence_items),

        "analyzed": len(results),

        "skipped": skipped,

        "results": results,

        "model_notes": model_notes,

    }

def analyze_case(case_id, actor):

    """Backward-compatible rule-based analysis for the existing UI."""

    evidence_items = get_evidence(case_id)

    if not evidence_items:

        raise ValueError("No evidence has been acquired for this case.")

    results = []

    for evidence in evidence_items:

        finding = inspect_evidence(evidence)

        save_analysis_result(

            evidence_id=evidence["evidence_id"],

            findings=finding,

            actor=actor,

        )

        results.append({

            "Evidence ID": evidence["evidence_id"],

            "File": evidence.get("file_name", ""),

            "Severity": finding["severity"],

            "Risk score": finding["risk_score"],

            "Integrity": finding["integrity"],

            "Findings": "; ".join(finding["findings"]),

        })

    return results

def run_system_checks():

    """Collect basic local system security-posture information."""

    checks = []

    operating_system = platform.system()

    checks.append({

        "Check": "Operating system",

        "Status": platform.platform(),

        "Severity": "INFO",

        "Details": "Platform information reported by Python.",

    })

    checks.append({

        "Check": "Python runtime",

        "Status": platform.python_version(),

        "Severity": "INFO",

        "Details": "Runtime version; not a complete patch assessment.",

    })

    checks.append({

        "Check": "Machine-learning dependencies",

        "Status": (

            "Available — scikit-learn"

            if SKLEARN_AVAILABLE

            else "Fallback available — pure-Python demo models"

        ),

        "Severity": "INFO" if SKLEARN_AVAILABLE else "REVIEW",

        "Details": (

            "NumPy and scikit-learn imported successfully."

            if SKLEARN_AVAILABLE

            else (

                "NumPy/scikit-learn are blocked or unavailable, so DataLens uses "

                "dependency-free demonstration Random Forest and Isolation Forest implementations. "

                "The predictions are experimental, not validated forensic conclusions. Original error: "

                f"{SKLEARN_IMPORT_ERROR or 'unknown error'}"

            )

        ),

    })

    if operating_system == "Windows":

        commands = [

            (

                "Windows Firewall",

                [

                    "powershell", "-NoProfile", "-Command",

                    "Get-NetFirewallProfile | Select-Object -ExpandProperty Enabled",

                ],

            ),

            (

                "Microsoft Defender real-time protection",

                [

                    "powershell", "-NoProfile", "-Command",

                    "(Get-MpComputerStatus).RealTimeProtectionEnabled",

                ],

            ),

        ]

        for label, command in commands:

            try:

                result = subprocess.run(

                    command, capture_output=True, text=True, timeout=10, check=False

                )

                output = result.stdout.strip()

                if result.returncode != 0 or not output:

                    status, severity = "Unable to verify", "UNKNOWN"

                elif label == "Windows Firewall":

                    values = [item.strip().lower() for item in output.splitlines() if item.strip()]

                    enabled = bool(values) and all(item == "true" for item in values)

                    status = "Enabled" if enabled else "One or more profiles disabled"

                    severity = "INFO" if enabled else "HIGH"

                else:

                    enabled = output.lower() == "true"

                    status = "Enabled" if enabled else "Disabled"

                    severity = "INFO" if enabled else "HIGH"

            except (OSError, subprocess.TimeoutExpired):

                status, severity = "Unable to verify", "UNKNOWN"

            checks.append({

                "Check": label,

                "Status": status,

                "Severity": severity,

                "Details": "Local system posture check.",

            })

    elif operating_system == "Linux":

        try:

            result = subprocess.run(

                ["systemctl", "is-active", "ufw"],

                capture_output=True, text=True, timeout=5, check=False,

            )

            ufw_status = result.stdout.strip() or "Unknown"

            checks.append({

                "Check": "UFW firewall service",

                "Status": ufw_status,

                "Severity": "INFO" if ufw_status == "active" else "REVIEW",

                "Details": (

                    "UFW status only. Other firewall frameworks, including nftables, may be configured."

                ),

            })

        except (OSError, subprocess.TimeoutExpired):

            checks.append({

                "Check": "UFW firewall service",

                "Status": "Unable to verify",

                "Severity": "UNKNOWN",

                "Details": "Could not query systemctl.",

            })

    return checks
